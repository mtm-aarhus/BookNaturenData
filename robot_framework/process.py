import hashlib
import hmac
import html
import io
import os
import re
import shutil
import tempfile
from datetime import date, datetime, timedelta
from urllib.parse import urlparse

import pandas as pd
import requests
from requests.auth import HTTPBasicAuth
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from office365.sharepoint.client_context import ClientContext

from OpenOrchestrator.orchestrator_connection.connection import OrchestratorConnection


def process(orchestrator_connection: OrchestratorConnection | None = None) -> None:

    # API
    basic_url = orchestrator_connection.get_constant('brugaarhusapiurl').value
    URL_BOOKINGER = f'{basic_url}/Booking/GetBookings'
    URL_STEDER = f'{basic_url}/Booking/GetPlaces'
    creds = orchestrator_connection.get_credential('brugaarhusapikey')
    USERNAME = creds.username
    PASSWORD = creds.password

    # SharePoint
    SHAREPOINT_SITE = orchestrator_connection.get_constant("PowerBiSharepoint").value.rstrip('/')
    BIBLIOTEK = "Delte dokumenter"   # internt URL-navn for "Delte dokumenter" - tjek adresselinjen i browseren
    SHAREPOINT_MAPPE = f"{urlparse(SHAREPOINT_SITE).path}/{BIBLIOTEK}/Data/BookNaturen"

    # Hemmelig nøgle til pseudonymisering - må ALDRIG ændres, ellers skifter alle Person-ID'er
    HASH_NØGLE = (orchestrator_connection.get_credential("brugaarhushashnoegle").password).encode()

    # ---------------------------------------------------------------------------
    # Bookinger
    # ---------------------------------------------------------------------------
    BOOKING_FILNAVN = "shelterbookinger.xlsx"
    BOOKING_ARK = "Bookinger"
    BOOKING_TABEL = "Bookinger"

    FØRSTE_DATO = date(2025, 1, 1)   # der findes ikke data fra før 2025
    DAGE_FREM = 730                  # bookinger ligger ofte ude i fremtiden
    DAGE_TILBAGE = 183               # daglig kørsel: ca. et halvt år tilbage

    INKLUDER_PERSONDATA = False

    BOOKING_KOLONNER = [
        "Booking-ID", "Oprettet", "Plads-ID", "Pladsnavn", "Starttidspunkt", "Antal timer",
        "Sluttidspunkt", "Antal personer", "Person-ID", "Aarhus-mail", "Bekræftet", "Aflyst",
        "Aflysningsdato", "Type", "Timer (dagsbooking)",
        # Pladsoplysninger
        "Plads OID", "Facilitetstype", "Facilitetstype-ID", "Arrangør-ID", "Breddegrad", "Længdegrad",
        "Plads slettet", "Plads oprettet", "Plads opdateret", "Kort beskrivelse", "Beskrivelse", "Billede-URL",
    ]
    if INKLUDER_PERSONDATA:
        BOOKING_KOLONNER += ["Fornavn", "Efternavn", "E-mail"]

    BOOKING_DATO_KOLONNER = ["Oprettet", "Starttidspunkt", "Sluttidspunkt", "Bekræftet", "Aflysningsdato",
                             "Plads oprettet", "Plads opdateret"]

    # ---------------------------------------------------------------------------
    # Faciliteter
    # ---------------------------------------------------------------------------
    STED_FILNAVN = "faciliteter.xlsx"
    STED_ARK = "Sheltere"
    STED_TABEL = "tblSheltere"
    INKLUDER_SLETTEDE_STEDER = True   # så gamle bookinger på nedlagte pladser stadig kan kobles til et sted

    STED_KOLONNER = [
        "Plads-ID", "Pladsnavn", "Facilitetstype", "Facilitetstype-ID", "Arrangør-ID",
        "Breddegrad", "Længdegrad", "Plads OID", "Slettet", "Oprettet", "Opdateret",
        "Kort beskrivelse", "Beskrivelse", "Billede-URL", "Billede-URL (lille)", "Antal billeder",
    ]
    STED_DATO_KOLONNER = ["Oprettet", "Opdateret"]

    # ---------------------------------------------------------------------------
    # Fælles hjælpefunktioner
    # ---------------------------------------------------------------------------
    def sharepoint_context() -> ClientContext:
        certification = orchestrator_connection.get_credential("SharePointCert")
        api = orchestrator_connection.get_credential("SharePointAPI")
        cert_credentials = {
            "tenant": api.username,
            "client_id": api.password,
            "thumbprint": certification.username,
            "cert_path": certification.password,
        }
        return ClientContext(SHAREPOINT_SITE).with_client_certificate(**cert_credentials)

    def parse_dato(værdi):
        """1900-datoer betyder 'ingen dato' og bliver til tomme celler."""
        if not værdi:
            return None
        dt = datetime.fromisoformat(værdi)
        return None if dt.year <= 1900 else dt

    def til_tal(værdi):
        try:
            return float(værdi)
        except (TypeError, ValueError):
            return None

    def html_til_tekst(værdi):
        """Fjerner HTML-tags fra beskrivelsen og laver den om til almindelig tekst."""
        if not værdi:
            return None
        tekst = re.sub(r"<br\s*/?>|</p>", "\n", værdi, flags=re.IGNORECASE)
        tekst = html.unescape(re.sub(r"<[^>]+>", "", tekst))
        tekst = re.sub(r"\n\s*\n+", "\n", tekst).strip()
        return tekst[:32000]   # Excel tillader max 32.767 tegn i en celle

    def lav_excel(df: pd.DataFrame, ark_navn: str, tabel_navn: str, dato_kolonner: list) -> bytes:
        """Bygger en formateret Excel-tabel i hukommelsen og returnerer den som bytes."""
        # Fjern usynlige kontroltegn, som Excel ikke tillader
        df = df.apply(lambda kol: kol.map(lambda v: ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v))

        rå = io.BytesIO()
        df.to_excel(rå, sheet_name=ark_navn, index=False)
        rå.seek(0)

        wb = load_workbook(rå)
        ws = wb[ark_navn]

        sidste_kolonne = get_column_letter(ws.max_column)
        tabel = Table(displayName=tabel_navn, ref=f"A1:{sidste_kolonne}{max(ws.max_row, 2)}")
        tabel.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tabel)

        for celle in ws[1]:
            celle.font = Font(name="Arial", bold=True, color="FFFFFF")
            celle.alignment = Alignment(vertical="center")

        for idx, kolonne in enumerate(df.columns, start=1):
            bogstav = get_column_letter(idx)
            if kolonne in dato_kolonner:
                for celle in ws[bogstav][1:]:
                    celle.number_format = "dd-mm-yyyy hh:mm"
                bredde = 18
            else:
                længste = max([len(str(kolonne))] + [len(str(v)) for v in df[kolonne].head(500)])
                bredde = min(max(længste + 2, 10), 45)
            ws.column_dimensions[bogstav].width = bredde
            for celle in ws[bogstav][1:]:
                celle.font = Font(name="Arial")

        ws.freeze_panes = "A2"

        færdig = io.BytesIO()
        wb.save(færdig)
        return færdig.getvalue()

    def upload_fil(ctx: ClientContext, filnavn: str, indhold: bytes):
        """Uploader i bidder via en lukket midlertidig fil (undgår Windows' låsefejl i office365)."""
        mappe_tmp = tempfile.mkdtemp()
        sti = os.path.join(mappe_tmp, filnavn)
        try:
            with open(sti, "wb") as f:
                f.write(indhold)
            mappe = ctx.web.get_folder_by_server_relative_url(SHAREPOINT_MAPPE)
            mappe.files.create_upload_session(sti, chunk_size=4 * 1024 * 1024, file_name=filnavn).execute_query()
        finally:
            shutil.rmtree(mappe_tmp, ignore_errors=True)

    # ---------------------------------------------------------------------------
    # Hjælpefunktioner - faciliteter
    # ---------------------------------------------------------------------------
    def hent_steder() -> list:
        payload = {
            "ArrangerID": 0,
            "IncludeDeleted": INKLUDER_SLETTEDE_STEDER,
        }
        response = requests.post(URL_STEDER, auth=HTTPBasicAuth(USERNAME, PASSWORD), json=payload, timeout=120)
        response.raise_for_status()
        data = response.json()
        return data if isinstance(data, list) else [data]

    def sted_til_række(s: dict) -> dict:
        billeder = s.get("Images") or {}
        store = billeder.get("Size1600") or []
        små = billeder.get("Size360") or []
        return {
            "Plads-ID": s.get("ID"),
            "Pladsnavn": s.get("Title"),
            "Facilitetstype": s.get("FacilityType"),
            "Facilitetstype-ID": s.get("FacilityTypeID"),
            "Arrangør-ID": s.get("ArrangerID"),
            "Breddegrad": til_tal(s.get("Latitude")),
            "Længdegrad": til_tal(s.get("Longitude")),
            "Plads OID": s.get("OID"),
            "Slettet": "Ja" if s.get("Deleted") else "Nej",
            "Oprettet": parse_dato(s.get("DateCreated")),
            "Opdateret": parse_dato(s.get("DateUpdated")),
            "Kort beskrivelse": s.get("ShortDescription"),
            "Beskrivelse": html_til_tekst(s.get("Description")),
            "Billede-URL": store[0] if store else None,
            "Billede-URL (lille)": små[0] if små else None,
            "Antal billeder": len(store),
        }

    # ---------------------------------------------------------------------------
    # Hjælpefunktioner - bookinger
    # ---------------------------------------------------------------------------
    def hent_bookinger(fra: date, til: date) -> list:
        bookinger = []
        for år in range(fra.year, til.year + 1):
            periode_fra, periode_til = max(fra, date(år, 1, 1)), min(til, date(år, 12, 31))
            payload = {
                "DateFrom": periode_fra.strftime("%Y%m%d"),
                "DateTo": periode_til.strftime("%Y%m%d"),
                "OnlyNoneCancelled": False,
                "ShowPlaceData": True,
            }
            orchestrator_connection.log_info(f"Henter bookinger {payload['DateFrom']}-{payload['DateTo']}")
            response = requests.post(URL_BOOKINGER, auth=HTTPBasicAuth(USERNAME, PASSWORD), json=payload, timeout=120)
            response.raise_for_status()
            data = response.json()
            bookinger.extend(data if isinstance(data, list) else [data])
        return bookinger

    def person_id(b: dict):
        """Fast, anonymt tal ud fra e-mail (eller navn, hvis e-mail mangler)."""
        nøgle = (b.get("Email") or "").strip().lower()
        if not nøgle:
            nøgle = f"{b.get('FirstName') or ''} {b.get('LastName') or ''}".strip().lower()
        if not nøgle:
            return None
        digest = hmac.new(HASH_NØGLE, nøgle.encode(), hashlib.sha256).hexdigest()
        return int(digest[:12], 16)

    def booking_til_række(b: dict) -> dict:
        start = parse_dato(b.get("Date"))
        timer = b.get("Nights") or 0          # "Nights" er reelt antal timer
        aflyst_dato = parse_dato(b.get("CancelDate"))
        bekræftet = parse_dato(b.get("ConfirmDate"))
        plads = b.get("PlaceData") or {}

        række = {
            "Booking-ID": b.get("ID"),
            "Oprettet": parse_dato(b.get("DateCreated")),
            "Plads-ID": b.get("PlaceID"),
            "Pladsnavn": plads.get("Title"),
            "Starttidspunkt": start,
            "Antal timer": timer,
            "Sluttidspunkt": start + timedelta(hours=timer) if start else None,
            "Antal personer": b.get("PeopleAmount"),
            "Person-ID": person_id(b),
            "Aarhus-mail": "Ja" if (b.get("Email") or "").strip().lower().endswith("@aarhus.dk") else "Nej",
            "Bekræftet": bekræftet,
            "Aflyst": "Ja" if aflyst_dato else "Nej",
            "Aflysningsdato": aflyst_dato,
            "Type": b.get("Type"),
            "Timer (dagsbooking)": b.get("Hours"),
            "Plads OID": plads.get("OID"),
            "Facilitetstype": plads.get("FacilityType"),
            "Facilitetstype-ID": plads.get("FacilityTypeID"),
            "Arrangør-ID": plads.get("ArrangerID"),
            "Breddegrad": til_tal(plads.get("Latitude")),
            "Længdegrad": til_tal(plads.get("Longitude")),
            "Plads slettet": "Ja" if plads.get("Deleted") else "Nej",
            "Plads oprettet": parse_dato(plads.get("DateCreated")),
            "Plads opdateret": parse_dato(plads.get("DateUpdated")),
            "Kort beskrivelse": plads.get("ShortDescription"),
            "Beskrivelse": html_til_tekst(plads.get("Description")),
            "Billede-URL": next(iter((plads.get("Images") or {}).get("Size1600") or []), None),
        }
        if INKLUDER_PERSONDATA:
            række["Fornavn"] = b.get("FirstName")
            række["Efternavn"] = b.get("LastName")
            række["E-mail"] = b.get("Email")
        return række

    def sorter_og_rens(df: pd.DataFrame) -> pd.DataFrame:
        return (
            df.drop_duplicates(subset="Booking-ID", keep="last")
            .sort_values(["Starttidspunkt", "Booking-ID"])
            .reset_index(drop=True)
        )

    # ---------------------------------------------------------------------------
    # Kørsel
    # ---------------------------------------------------------------------------
    ctx = sharepoint_context()

    # 1. Faciliteter - hele listen hentes og overskriver filen hver gang
    steder = hent_steder()
    df_steder = (
        pd.DataFrame([sted_til_række(s) for s in steder], columns=STED_KOLONNER)
        .drop_duplicates(subset="Plads-ID", keep="last")
        .sort_values(["Facilitetstype", "Pladsnavn"])
        .reset_index(drop=True)
    )
    upload_fil(ctx, STED_FILNAVN, lav_excel(df_steder, STED_ARK, STED_TABEL, STED_DATO_KOLONNER))
    orchestrator_connection.log_info(f"Uploadet {STED_FILNAVN} med {len(df_steder)} faciliteter")

    # 2. Bookinger - hent eksisterende fil fra SharePoint
    download = io.BytesIO()
    ctx.web.get_file_by_server_relative_url(f"{SHAREPOINT_MAPPE}/{BOOKING_FILNAVN}").download(download).execute_query()
    download.seek(0)
    eksisterende = pd.read_excel(download, sheet_name=BOOKING_ARK)

    manglende = [k for k in BOOKING_KOLONNER if k not in eksisterende.columns]
    if manglende:
        raise ValueError(
            f"Arket mangler kolonnerne {manglende}. Kør første kørsel igen, eller tjek at "
            f"INKLUDER_PERSONDATA er sat ens i begge scripts."
        )
    eksisterende = eksisterende[BOOKING_KOLONNER]
    orchestrator_connection.log_info(f"Hentet {BOOKING_FILNAVN} fra SharePoint med {len(eksisterende)} bookinger")

    # Hent det seneste halve år (og frem) fra API'et
    fra_dato = max(FØRSTE_DATO, date.today() - timedelta(days=DAGE_TILBAGE))
    til_dato = date.today() + timedelta(days=DAGE_FREM)
    orchestrator_connection.log_info(f"Daglig kørsel: henter {fra_dato} til {til_dato}")
    nye = pd.DataFrame([booking_til_række(b) for b in hent_bookinger(fra_dato, til_dato)], columns=BOOKING_KOLONNER)

    # Nye data vinder: eksisterende bookinger overskrives (fx nye aflysninger), nye tilføjes
    samlet = sorter_og_rens(pd.concat([eksisterende, nye], ignore_index=True))

    upload_fil(ctx, BOOKING_FILNAVN, lav_excel(samlet, BOOKING_ARK, BOOKING_TABEL, BOOKING_DATO_KOLONNER))
    orchestrator_connection.log_info(
        f"Uploadet {BOOKING_FILNAVN}: {len(samlet)} bookinger i alt ({len(nye)} hentet i denne kørsel)"
    )